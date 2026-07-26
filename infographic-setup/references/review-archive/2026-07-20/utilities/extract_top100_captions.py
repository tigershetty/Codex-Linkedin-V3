from __future__ import annotations

import argparse
import re
from pathlib import Path

from openpyxl import load_workbook


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("start", type=int)
    parser.add_argument("end", type=int)
    args = parser.parse_args()

    workbook = Path("infographic-setup/references/top 100/Reference File and Caption.xlsx")
    sheet = load_workbook(workbook, read_only=True, data_only=True).active
    rows = list(sheet.iter_rows(min_row=2, values_only=True))

    for index, row in enumerate(rows, 1):
        if not args.start <= index <= args.end:
            continue
        ref = row[0]
        caption = str(row[1] or "").strip()
        lines = [line.strip() for line in caption.splitlines() if line.strip()]
        urls = re.findall(r"https?://\S+", caption)
        opener = lines[0][:240] if lines else ""
        closer = " | ".join(lines[-7:])[:1_000]
        print(f"REF {index} | {ref}")
        print(f"OPEN: {opener}")
        print(f"CLOSE: {closer}")
        print(f"URLS: {' '.join(urls)[:800]}")
        print()


if __name__ == "__main__":
    main()
