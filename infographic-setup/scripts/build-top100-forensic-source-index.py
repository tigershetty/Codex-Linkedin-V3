#!/usr/bin/env python3
"""Extract Top-100 caption context without interpreting visual craft."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "references/top 100/Reference File and Caption.xlsx"
OUTPUT = ROOT / "references/creative-review/top100-forensic-source-context.jsonl"


def first_nonempty_line(text: str) -> str:
    for line in text.splitlines():
        cleaned = line.strip().lstrip("•-→👉 ")
        if cleaned:
            return cleaned
    return ""


def main() -> None:
    sheet = load_workbook(WORKBOOK, read_only=True, data_only=True).active
    records = []
    for row in sheet.iter_rows(min_row=3, values_only=True):
        number, caption = row[0], row[1]
        if not isinstance(number, int):
            continue
        text = str(caption or "").strip()
        records.append(
            {
                "top100_asset_number": number,
                "caption_available": bool(text),
                "caption_source": "references/top 100/Reference File and Caption.xlsx:Sheet1",
                "caption_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest() if text else None,
                "caption_hook": first_nonempty_line(text),
                "caption_text": text,
                "extraction_note": "Raw source context only. It does not establish visual mechanics or factual validity."
            }
        )
    records.sort(key=lambda item: item["top100_asset_number"])
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text("".join(json.dumps(record, ensure_ascii=False) + "\n" for record in records), encoding="utf-8")
    print(f"Wrote {len(records)} source-context records to {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
